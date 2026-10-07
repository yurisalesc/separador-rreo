"""Exportação de PDFs por anexo, planilha de conferência e ZIP."""

from __future__ import annotations

import zipfile
from collections.abc import Iterable, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from separador_rreo.constants import ANEXOS_PADRAO
from separador_rreo.models import AnalysisResult, ExportResult, Finding
from separador_rreo.pdf_source import PdfDocument
from separador_rreo.text import safe_dirname


class RreoExporter:
    """
    Gera artefatos a partir de um `AnalysisResult`.

    Responsabilidade única: persistir PDFs, planilha e ZIP.
    """

    def __init__(self, allowed_annexes: Sequence[int] = ANEXOS_PADRAO) -> None:
        self._annexes = tuple(allowed_annexes)

    def export(
        self,
        document: PdfDocument,
        analysis: AnalysisResult,
        destination: Path | str,
    ) -> ExportResult:
        output_dir = Path(destination)
        output_dir.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Conferencia RREO"
        sheet.append(
            [
                "Prefeitura",
                "Anexo",
                "Pagina inicial FEMURN",
                "Pagina final FEMURN",
                "Total de paginas",
                "Situacao",
                "Conferencia manual",
                "Titulo encontrado",
                "Arquivo",
            ]
        )

        rows: list[dict[str, object]] = []
        pdf_count = 0

        for municipality, findings in analysis.findings_by_municipality.items():
            folder = output_dir / safe_dirname(municipality)
            folder.mkdir(exist_ok=True)

            for annex in self._annexes:
                groups = _merge_sequential(
                    sorted(
                        (item for item in findings if item.annex == annex),
                        key=lambda item: (item.page_start, item.page_end),
                    )
                )

                if not groups:
                    sheet.append(
                        [
                            municipality,
                            annex,
                            None,
                            None,
                            None,
                            "Nao localizado",
                            "Verificar",
                            None,
                            None,
                        ]
                    )
                    rows.append(
                        {
                            "municipality": municipality,
                            "annex": annex,
                            "status": "Nao localizado",
                            "file": None,
                        }
                    )
                    continue

                for publication_index, finding in enumerate(groups, start=1):
                    suffix = (
                        f"_publicacao_{publication_index}" if len(groups) > 1 else ""
                    )
                    filename = f"RREO_{annex:02d}_{municipality}{suffix}.pdf"
                    document.extract_pages(
                        finding.page_start,
                        finding.page_end,
                        folder / filename,
                    )
                    pdf_count += 1

                    label = finding.status
                    if len(groups) > 1:
                        label = f"{finding.status} - multiplas publicacoes"

                    sheet.append(
                        [
                            municipality,
                            annex,
                            finding.page_start,
                            finding.page_end,
                            finding.page_end - finding.page_start + 1,
                            label,
                            "Pendente",
                            finding.title,
                            filename,
                        ]
                    )
                    rows.append(
                        {
                            "municipality": municipality,
                            "annex": annex,
                            "page_start": finding.page_start,
                            "page_end": finding.page_end,
                            "status": label,
                            "file": filename,
                        }
                    )

        _style_sheet(sheet)
        spreadsheet = output_dir / "Controle_Publicacoes.xlsx"
        workbook.save(str(spreadsheet))

        zip_path = output_dir.with_suffix(".zip")
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in output_dir.rglob("*"):
                if path.is_file():
                    archive.write(path, path.relative_to(output_dir))

        return ExportResult(
            output_dir=output_dir.resolve(),
            spreadsheet=spreadsheet.resolve(),
            zip_path=zip_path.resolve(),
            pdf_count=pdf_count,
            rows=tuple(rows),
        )


def _merge_sequential(findings: Iterable[Finding]) -> list[Finding]:
    """Unifica partes sequenciais do mesmo anexo em um único PDF completo."""
    groups: list[Finding] = []
    for finding in findings:
        if groups and finding.page_start <= groups[-1].page_end + 1:
            previous = groups[-1]
            if previous.status == "Encontrado" and finding.status == "Encontrado":
                status = "Encontrado - partes reunidas"
            else:
                status = "Verificar partes reunidas"
            groups[-1] = Finding(
                annex=finding.annex,
                page_start=previous.page_start,
                page_end=max(previous.page_end, finding.page_end),
                status=status,
                title=f"{previous.title} / {finding.title}",
            )
        else:
            groups.append(finding)
    return groups


def _style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{sheet.max_row}"
    header_fill = PatternFill("solid", fgColor="164E63")
    header_font = Font(bold=True, color="FFFFFF")
    header_align = Alignment(wrap_text=True, vertical="center")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
    sheet.row_dimensions[1].height = 32

    widths = {
        "A": 27,
        "B": 9,
        "C": 23,
        "D": 23,
        "E": 17,
        "F": 39,
        "G": 21,
        "H": 93,
        "I": 51,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    validation = DataValidation(type="list", formula1='"Pendente,Correto,Verificar"')
    sheet.add_data_validation(validation)
    validation.add(f"G2:G{sheet.max_row}")

    for row in sheet.iter_rows(min_row=2):
        cell = row[5]
        value = str(cell.value or "")
        if value == "Encontrado":
            color = "DCFCE7"
        elif "Verificar" in value or "multiplas" in value:
            color = "FEF3C7"
        else:
            color = "FEE2E2"
        cell.fill = PatternFill("solid", fgColor=color)
