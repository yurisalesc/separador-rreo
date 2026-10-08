"""Geração da planilha de conferência RREO."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from separador_rreo.enumeracoes import SituacaoConferencia, SituacaoOcorrencia
from separador_rreo.modelos import LinhaConferencia


class GeradorPlanilhaConferencia:
    """Cria a planilha e concentra todas as decisões de apresentação."""

    CABECALHOS = (
        "Prefeitura",
        "Anexo",
        "Pagina inicial FEMURN",
        "Pagina final FEMURN",
        "Total de paginas",
        "Situacao",
        "Conferencia manual",
        "Titulo encontrado",
        "Arquivo",
    )
    LARGURAS = {
        "A": 27,
        "B": 9,
        "C": 23,
        "D": 23,
        "E": 17,
        "F": 39,
        "G": 21,
        "H": 93,
        "I": 51,
    }

    def gerar(self, linhas: Sequence[LinhaConferencia], destino: Path) -> None:
        pasta = destino.parent
        pasta.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Conferencia RREO"
        planilha.append(list(self.CABECALHOS))

        for linha in linhas:
            planilha.append(
                [
                    linha.municipio,
                    int(linha.anexo),
                    linha.pagina_inicial,
                    linha.pagina_final,
                    linha.total_paginas,
                    linha.rotulo_situacao,
                    linha.conferencia.value,
                    linha.titulo,
                    linha.arquivo,
                ]
            )

        self._estilizar(planilha)
        workbook.save(str(destino))

    def _estilizar(self, planilha) -> None:
        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = f"A1:I{planilha.max_row}"

        preenchimento = PatternFill("solid", fgColor="164E63")
        fonte = Font(bold=True, color="FFFFFF")
        alinhamento = Alignment(wrap_text=True, vertical="center")
        for celula in planilha[1]:
            celula.fill = preenchimento
            celula.font = fonte
            celula.alignment = alinhamento
        planilha.row_dimensions[1].height = 32

        for coluna, largura in self.LARGURAS.items():
            planilha.column_dimensions[coluna].width = largura

        opcoes = ",".join(item.value for item in SituacaoConferencia)
        validacao = DataValidation(type="list", formula1=f'"{opcoes}"')
        planilha.add_data_validation(validacao)
        validacao.add(f"G2:G{planilha.max_row}")

        for linha in planilha.iter_rows(min_row=2):
            celula = linha[5]
            valor = str(celula.value or "")
            celula.fill = PatternFill("solid", fgColor=self._cor_da_situacao(valor))

    @staticmethod
    def _cor_da_situacao(valor: str) -> str:
        if valor == SituacaoOcorrencia.ENCONTRADO.value:
            return "DCFCE7"
        if "Verificar" in valor or "multiplas" in valor:
            return "FEF3C7"
        return "FEE2E2"
